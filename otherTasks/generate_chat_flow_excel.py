import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Test Cases'

headers = [
    'project_id', 'suite_id', 'unique_id', 'title', 'decription',
    'precondition', 'priority', 'type', 'status', 'tags',
    'steps', 'step_desc', 'expected_desc', 'is_automated', 'user_story'
]
ws.append(headers)

col_widths = {
    'A': 12.0, 'B': 12.0, 'C': 13.0, 'D': 45.0, 'E': 45.0,
    'F': 45.0, 'G': 15.0, 'H': 12.0, 'I': 12.0, 'J': 12.0,
    'K': 12.0, 'L': 45.0, 'M': 45.0, 'N': 14.0, 'O': 45.0
}
for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

header_fill = PatternFill(start_color='001F4E78', end_color='001F4E78', fill_type='solid')
header_font = Font(name='Calibri', size=11, bold=True, color='00FFFFFF')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

for col_idx in range(1, 16):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align

thin_border = Border(
    left=Side(style='thin', color='00D9D9D9'),
    right=Side(style='thin', color='00D9D9D9'),
    top=Side(style='thin', color='00D9D9D9'),
    bottom=Side(style='thin', color='00D9D9D9')
)

test_cases = [
    # --- PRECONDITION ---
    {
        'unique_id': 'TC_PRE_001',
        'title': 'Verifikasi autentikasi akun & langganan aktif Everpro Chat',
        'decription': 'Memastikan user yang memiliki akun valid dan paket langganan aktif dapat mengakses modul Everpro Chat dan seluruh sub-fiturnya.',
        'precondition': '1. Akun pengguna terdaftar di Everpro.\n2. Akun memiliki subscription paket Everpro Chat yang aktif.\n3. Minimal 1 nomor WABA terhubung ke akun.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Login ke dashboard Everpro.\n2. Masuk ke menu Everpro Chat.\n3. Periksa aksesibilitas menu Pengaturan Alur Chat dan Chat Room.',
        'expected_desc': '1. Pengguna berhasil masuk ke dashboard Everpro Chat tanpa peringatan subscription kadaluwarsa.\n2. Seluruh modul utama tampil dan dapat diakses sesuai role.',
        'is_automated': 'FALSE',
        'user_story': 'Global Pre-condition'
    },
    {
        'unique_id': 'TC_PRE_002',
        'title': 'Verifikasi pencegahan akses bagi akun dengan subscription expired / inaktif',
        'decription': 'Memastikan user dengan status subscription expired/inactive diblokir dari akses fitur Everpro Chat.',
        'precondition': '1. Akun terdaftar di Everpro.\n2. Status subscription Everpro Chat adalah "Inactive" / "Expired".',
        'priority': 'P1 - Critical',
        'step_desc': '1. Login ke dashboard Everpro menggunakan akun expired.\n2. Buka menu Everpro Chat atau akses direct URL ke /pengaturan-alur-chat.',
        'expected_desc': '1. Tampil banner/modal bahwa paket subscription tidak aktif.\n2. Akses ke menu chat diblokir dan muncul CTA untuk perpanjangan paket.',
        'is_automated': 'FALSE',
        'user_story': 'Global Pre-condition'
    },

    # --- RBAC ACCESS CONTROL ---
    {
        'unique_id': 'TC_RBAC_001',
        'title': 'Verifikasi full access role Owner pada menu Pengaturan Alur Chat & AI Agent',
        'decription': 'Memastikan role Owner dapat mengakses, membuat, mengedit, mengaktifkan/menonaktifkan, dan menghapus alur chat serta mengonfigurasi AI Agent.',
        'precondition': '1. Login sebagai akun dengan role Owner.\n2. Terdapat nomor WABA aktif pada akun.',
        'priority': 'P2 - High',
        'step_desc': '1. Navigasi ke menu Pengaturan Alur Chat.\n2. Buat, edit, aktifkan toggle, dan hapus salah satu flow.\n3. Navigasi ke menu AI Agent dan lakukan konfigurasi.',
        'expected_desc': '1. Menu Pengaturan Alur Chat dan AI Agent terbuka sempurna.\n2. Semua aksi CRUD (Create, Read, Update, Delete) dan toggle alur chat serta AI Agent berhasil dijalankan.',
        'is_automated': 'FALSE',
        'user_story': 'Role-Based Access Control (RBAC)'
    },
    {
        'unique_id': 'TC_RBAC_002',
        'title': 'Verifikasi full access role Supervisor pada menu Pengaturan Alur Chat & AI Agent',
        'decription': 'Memastikan role Supervisor memiliki kewenangan manajemen penuh terhadap alur chat dan AI Agent setara dengan Owner.',
        'precondition': '1. Login sebagai akun dengan role Supervisor.\n2. Terdapat nomor WABA aktif pada akun.',
        'priority': 'P2 - High',
        'step_desc': '1. Masuk ke menu Pengaturan Alur Chat.\n2. Lakukan modifikasi alur chat dan simpan perubahan.\n3. Buka konfigurasi AI Agent dan simpan aturan transfer/fallback.',
        'expected_desc': '1. Supervisor dapat melihat list alur, membuka visual builder, dan menyimpan perubahan flow.\n2. Supervisor dapat mengelola konfigurasi AI Agent tanpa batasan otorisasi.',
        'is_automated': 'FALSE',
        'user_story': 'Role-Based Access Control (RBAC)'
    },
    {
        'unique_id': 'TC_RBAC_003',
        'title': 'Verifikasi pembatasan akses role CS Agent terhadap menu Pengaturan Alur Chat (Hidden / Blocked)',
        'decription': 'Memastikan role CS (Closing/Operational) tidak dapat melihat atau mengakses menu Pengaturan Alur Chat maupun konfigurasi AI Agent.',
        'precondition': '1. Login sebagai akun dengan role CS Agent (Closing / Operational).',
        'priority': 'P2 - High',
        'step_desc': '1. Periksa sidebar menu navigasi Everpro Chat.\n2. Coba akses langsung via URL browser /pengaturan-alur-chat atau /ai-agent/config.',
        'expected_desc': '1. Menu Pengaturan Alur Chat dan AI Agent tersembunyi (hidden) dari sidebar menu.\n2. Akses direct URL menghasilkan halaman Error 403 Forbidden atau dialihkan ke halaman Chat Room.',
        'is_automated': 'FALSE',
        'user_story': 'Role-Based Access Control (RBAC)'
    },
    {
        'unique_id': 'TC_RBAC_004',
        'title': 'Verifikasi kewenangan CS Agent untuk melakukan aksi Ambil Alih Chat pada Chat Room',
        'decription': 'Memastikan CS Agent yang ditugaskan pada suatu percakapan memiliki hak akses untuk menekan tombol Ambil Alih Chat (Turn off flow).',
        'precondition': '1. Login sebagai akun dengan role CS Agent.\n2. Terdapat percakapan masuk di Chat Room yang sedang ditangani oleh AI Agent atau Template.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Buka halaman Chat Room.\n2. Pilih room percakapan aktif yang sedang ditangani AI Agent.\n3. Klik tombol overlay Ambil Alih Chat.',
        'expected_desc': '1. Tombol Ambil Alih Chat dapat diklik oleh CS Agent.\n2. Flow berhasil di-turn off dan CS Agent dapat langsung mengetik serta mengirim pesan balasan manual.',
        'is_automated': 'FALSE',
        'user_story': 'Role-Based Access Control (RBAC)'
    },

    # --- MODUL 1: Pengaturan Alur Chat List & Column Updates (US 1) ---
    {
        'unique_id': 'TC_FLOW_001',
        'title': 'Verifikasi perubahan nama fitur dari "Automasi" menjadi "Pengaturan Alur Chat"',
        'decription': 'Memastikan rebranding nama menu dan judul halaman telah diperbarui dari Automasi menjadi Pengaturan Alur Chat.',
        'precondition': '1. User login sebagai Owner/Supervisor.\n2. Masuk ke dashboard Everpro Chat.',
        'priority': 'P3 - Medium',
        'step_desc': '1. Periksa label menu pada sidebar navigasi.\n2. Klik menu Pengaturan Alur Chat.\n3. Periksa header / title pada halaman yang terbuka.',
        'expected_desc': '1. Label menu di sidebar menampilkan teks "Pengaturan Alur Chat" (menggantikan "Automasi").\n2. Judul halaman menampilkan "Pengaturan Alur Chat".',
        'is_automated': 'FALSE',
        'user_story': 'Modul 1: Pengaturan Alur Chat List & Structure (US 1)'
    },
    {
        'unique_id': 'TC_FLOW_002',
        'title': 'Verifikasi struktur dan kelengkapan kolom tabel pada list Pengaturan Alur Chat',
        'decription': 'Memastikan daftar alur menampilkan seluruh kolom baru yang dipersyaratkan: Name, Assigned to phone number, Lead source, Keyword source, First responders, Involve human, dan Active/inactive toggle.',
        'precondition': '1. User login sebagai Owner/Supervisor.\n2. Terdapat beberapa flow aktif/inaktif yang sudah terdaftar.',
        'priority': 'P2 - High',
        'step_desc': '1. Buka halaman Pengaturan Alur Chat.\n2. Periksa header kolom tabel data alur chat.\n3. Periksa rendering nilai data pada masing-masing baris.',
        'expected_desc': '1. Tabel memuat kolom: Name, Assigned to phone number, Lead source (CTWA/LPWA/Balasan broadcast/Organic), Keyword source (Default/custom), First responders (Human Agent/AI Agent/Template), Involve human from the start (checklist Y/N), dan Active/inactive (toggle Y/N).\n2. Seluruh data ter-render dengan rapi dan sesuai konfigurasi masing-masing alur.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 1: Pengaturan Alur Chat List & Structure (US 1)'
    },

    # --- MODUL 2: Primary Default Flow & Multi-trigger Logic (US 2) ---
    {
        'unique_id': 'TC_FLOW_003',
        'title': 'Verifikasi pembentukan otomatis Primary Default Flow saat onboarding nomor WABA baru',
        'decription': 'Memastikan sistem secara otomatis membuat Tim CS default flow dengan parameter default ketika nomor WABA baru ditambahkan ke akun.',
        'precondition': '1. User onboarding / menambahkan nomor WhatsApp baru (misal: "123").\n2. Belum ada konfigurasi flow lain pada nomor tersebut.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Selesaikan onboarding nomor baru "123".\n2. Buka menu Pengaturan Alur Chat.\n3. Periksa baris alur default untuk nomor "123".',
        'expected_desc': '1. Alur "Tim CS default flow" otomatis terbentuk dengan nilai:\n- Name: Tim CS default flow\n- Phone number: 123\n- Source: default\n- Keyword: default\n- First response: Human agent\n- Involve human from the start: yes\n- Active: yes\n- Response: transfer directly to "primary" division (divisi utama).\n2. Tombol edit/deaktivasi/hapus pada flow ini dalam keadaan disabled / protected.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 2: Primary Default Flow & Multi-trigger (US 2)'
    },
    {
        'unique_id': 'TC_FLOW_004',
        'title': 'Verifikasi proteksi Primary Default Flow dari tindakan edit, deaktivasi, dan hapus',
        'decription': 'Memastikan pengguna tidak dapat mengubah parameter, mematikan toggle aktif, atau menghapus Primary Default Flow secara langsung.',
        'precondition': '1. Buka list Pengaturan Alur Chat.\n2. Temukan baris Primary Default Flow ("Tim CS default flow" / "AI Agent default flow").',
        'priority': 'P1 - Critical',
        'step_desc': '1. Arahkan kursor ke aksi baris Primary Default Flow.\n2. Coba klik tombol Edit, Toggle Deactive, atau Hapus.',
        'expected_desc': '1. Toggle status terkunci (disabled) dan aksi Hapus/Edit tidak tersedia atau menampilkan tooltip informasi bahwa flow default sistem tidak dapat dihapus/diubah secara langsung.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 2: Primary Default Flow & Multi-trigger (US 2)'
    },
    {
        'unique_id': 'TC_FLOW_005',
        'title': 'Verifikasi override Primary Default Flow saat toggle AI Autoreply (agentic) diaktifkan',
        'decription': 'Memastikan Primary Default Flow otomatis ter-override menjadi AI Agent A default flow saat fitur AI Autoreply diaktifkan.',
        'precondition': '1. AI Agent A telah ditugaskan ke nomor WABA "123".\n2. Toggle "AI Auto-reply (agentic)" pada AI Agent A dalam kondisi OFF.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Buka konfigurasi AI Agent A.\n2. Aktifkan toggle "AI Auto-reply (agentic)" ke ON dan simpan.\n3. Buka menu Pengaturan Alur Chat dan periksa baris Primary Default Flow nomor "123".',
        'expected_desc': '1. Primary Default Flow otomatis ter-override menjadi:\n- Name: AI Agent A default flow\n- First response: AI Agent\n- Involve human from the start: no\n- Active: yes\n- Response: transfer directly to AI Agent A.\n2. Pesan masuk baru dengan source/keyword default langsung ditangani oleh AI Agent A.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 2: Primary Default Flow & Multi-trigger (US 2)'
    },
    {
        'unique_id': 'TC_FLOW_006',
        'title': 'Verifikasi pengembalian Primary Default Flow ke Tim CS saat toggle AI Autoreply dimatikan',
        'decription': 'Memastikan Primary Default Flow otomatis kembali ke "Tim CS default flow" saat toggle AI Autoreply dinonaktifkan.',
        'precondition': '1. Toggle "AI Auto-reply (agentic)" AI Agent A nomor "123" sebelumnya aktif (ON).',
        'priority': 'P1 - Critical',
        'step_desc': '1. Buka menu konfigurasi AI Agent A.\n2. Matikan toggle "AI Auto-reply (agentic)" ke OFF dan konfirmasi peringatan.\n3. Buka menu Pengaturan Alur Chat dan periksa baris Primary Default Flow nomor "123".',
        'expected_desc': '1. Primary Default Flow otomatis kembali menjadi "Tim CS default flow" (First response: Human agent, Involve human: yes, Target: divisi utama).\n2. Pesan masuk baru tanpa trigger spesifik kembali dialihkan ke Tim CS divisi utama.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 2: Primary Default Flow & Multi-trigger (US 2)'
    },
    {
        'unique_id': 'TC_FLOW_007',
        'title': 'Verifikasi penghapusan otomatis Primary Default Flow saat nomor WABA dihapus dari akun',
        'decription': 'Memastikan lifecycle flow default ikut terhapus bersih saat nomor WABA terkait dihapus dari akun Everpro.',
        'precondition': '1. Terdapat nomor WABA "123" yang memiliki Primary Default Flow.',
        'priority': 'P2 - High',
        'step_desc': '1. Masuk ke pengaturan WABA / integrasi nomor.\n2. Hapus nomor WABA "123" dari akun Everpro.\n3. Buka menu Pengaturan Alur Chat.',
        'expected_desc': '1. Seluruh entitas alur chat yang terikat pada nomor "123" termasuk Primary Default Flow ikut terhapus dari daftar.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 2: Primary Default Flow & Multi-trigger (US 2)'
    },
    {
        'unique_id': 'TC_FLOW_008',
        'title': 'Verifikasi pencegahan duplikasi kombinasi trigger pertama (Phone + Source + Keyword)',
        'decription': 'Memastikan sistem menolak penyimpanan alur baru jika kombinasi Phone Number, Source, dan Keyword sama persis dengan alur yang sudah ada.',
        'precondition': '1. Telah ada Flow A pada nomor "123" dengan Source "C" dan Keyword "PROMO".',
        'priority': 'P1 - Critical',
        'step_desc': '1. Klik tombol Buat Alur Chat Baru.\n2. Pilih nomor WABA "123", pilih Source "C", dan masukkan Keyword "PROMO" (atau "promo" - case-insensitive).\n3. Klik Simpan Alur.',
        'expected_desc': '1. Sistem mencegah penyimpanan alur.\n2. Muncul pesan error validasi: "flow with Phone number 123, source C, and Keyword PROMO is already exist".',
        'is_automated': 'FALSE',
        'user_story': 'Modul 2: Primary Default Flow & Multi-trigger (US 2)'
    },
    {
        'unique_id': 'TC_FLOW_009',
        'title': 'Verifikasi evaluasi prioritas trigger: Source over Keyword (Lead Source match > Keyword match)',
        'decription': 'Memastikan ketika pesan masuk memiliki Source B dan Keyword A, sistem mengeksekusi alur dengan Source B (Flow Y) bukan alur dengan Keyword A (Flow X).',
        'precondition': '1. Terdaftar alur pada nomor "123":\n- Flow Tim CS default (Source: default, Keyword: default)\n- Flow X (Source: default, Keyword: A)\n- Flow Y (Source: B, Keyword: default)\n- Flow Z (Source: C, Keyword: D)\n- Flow Z-default (Source: C, Keyword: default)',
        'priority': 'P1 - Critical',
        'step_desc': '1. Kirim pesan masuk WhatsApp ke nomor "123" yang berasal dari Lead Source B dengan isi teks "A".\n2. Periksa alur yang tereksekusi pada percakapan tersebut.',
        'expected_desc': '1. Sistem mengeksekusi Flow Y (karena Source B cocok secara spesifik, mengabaikan keyword match pada Flow X).\n2. Respon percakapan mengikuti aksi yang dikonfigurasi pada Flow Y.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 2: Primary Default Flow & Multi-trigger (US 2)'
    },
    {
        'unique_id': 'TC_FLOW_010',
        'title': 'Verifikasi pencocokan exact match & case-insensitive pada custom keyword alur chat',
        'decription': 'Memastikan keyword dipicu hanya pada kata yang cocok persis (exact match) dan mengabaikan huruf besar/kecil (case-insensitive).',
        'precondition': '1. Terdaftar alur Flow K dengan Source default dan Keyword kustom "INFO DISKON".',
        'priority': 'P1 - Critical',
        'step_desc': '1. Kirim pesan masuk dengan teks "info diskon" (lowercase exact).\n2. Kirim pesan masuk dengan teks "INFO DISKON" (uppercase exact).\n3. Kirim pesan masuk dengan teks "halo saya mau tanya info diskon dong" (partial / non-exact).',
        'expected_desc': '1. Pesan 1 dan 2 berhasil memicu Flow K (case-insensitive exact match).\n2. Pesan 3 tidak memicu Flow K (karena exact match) dan diteruskan ke default fallback flow.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 2: Primary Default Flow & Multi-trigger (US 2)'
    },
    {
        'unique_id': 'TC_FLOW_011',
        'title': 'Verifikasi pembuatan otomatis default keyword pair saat membuat flow custom keyword',
        'decription': 'Memastikan saat user membuat flow Z dengan Source C dan Keyword D, sistem otomatis membuat alur pendamping "Z - default keyword" (Source: C, Keyword: default).',
        'precondition': '1. User membuka form Buat Alur Chat pada nomor "123".',
        'priority': 'P1 - Critical',
        'step_desc': '1. Pilih Source "C".\n2. Aktifkan toggle "Gunakan Keyword" dan masukkan keyword "D".\n3. Konfigurasi respon dan klik Simpan.',
        'expected_desc': '1. Alur "Z" berhasil dibuat.\n2. Sistem secara otomatis membuat entitas alur "Z - default keyword" dengan nilai:\n- Name: Z - default keyword\n- Phone number: 123, Source: C, Keyword: default\n- First response: Human agent, Involve human: yes, Active: yes, Target: divisi utama.\n3. Parameter Name, Phone number, Source, Keyword, dan Status Active pada flow default pair terkunci (read-only), namun Response dapat dimodifikasi di visual builder.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 2: Primary Default Flow & Multi-trigger (US 2)'
    },
    {
        'unique_id': 'TC_FLOW_012',
        'title': 'Verifikasi eksekusi alur default keyword pair pada pesan dengan keyword tak cocok',
        'decription': 'Memastikan pesan masuk dari Source C dengan keyword yang tidak terdaftar (Keyword E) dialihkan ke alur "Z - default keyword".',
        'precondition': '1. Alur Flow Z (Source C, Keyword D) dan Flow Z-default (Source C, Keyword default) aktif pada nomor "123".',
        'priority': 'P1 - Critical',
        'step_desc': '1. Kirim pesan masuk WhatsApp dari Source C dengan isi teks "E".\n2. Periksa alur yang dijalankan oleh sistem.',
        'expected_desc': '1. Sistem mengeksekusi alur "Z - default keyword".\n2. Percakapan diteruskan langsung ke Tim CS divisi utama sesuai konfigurasi flow default pair.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 2: Primary Default Flow & Multi-trigger (US 2)'
    },
    {
        'unique_id': 'TC_FLOW_013',
        'title': 'Verifikasi sinkronisasi status aktif/inaktif dan penghapusan otomatis flow default keyword pair',
        'decription': 'Memastikan status active/inactive alur default pair otomatis tersinkronisasi dengan alur induknya, dan ikut terhapus saat alur induk dihapus atau toggle keyword dimatikan.',
        'precondition': '1. Terdapat Flow Z (custom keyword) dan pasangan otomatisnya Flow Z-default.',
        'priority': 'P2 - High',
        'step_desc': '1. Ubah status Flow Z menjadi Inactive dan simpan. Periksa status Flow Z-default.\n2. Aktifkan kembali Flow Z. Periksa status Flow Z-default.\n3. Hapus Flow Z (atau matikan toggle custom keyword pada Flow Z).',
        'expected_desc': '1. Saat Flow Z Inactive, Flow Z-default otomatis menjadi Inactive.\n2. Saat Flow Z Active, Flow Z-default otomatis menjadi Active kembali.\n3. Saat Flow Z dihapus atau toggle keyword dimatikan, alur Flow Z-default otomatis ikut terhapus.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 2: Primary Default Flow & Multi-trigger (US 2)'
    },

    # --- MODUL 3: Quick Setup & Activation of AI Agent (US 3) ---
    {
        'unique_id': 'TC_AI_001',
        'title': 'Verifikasi navigasi form konfigurasi AI Agent (Tab Basic Info, Knowledge, Style)',
        'decription': 'Memastikan form konfigurasi AI Agent dapat diakses dan menampilkan tab Basic Info, Knowledge, dan Communication Style secara lengkap.',
        'precondition': '1. User login sebagai Owner/Supervisor.\n2. Masuk ke menu AI Agent (Daftar Agen).',
        'priority': 'P2 - High',
        'step_desc': '1. Klik tombol "Konfigurasi" pada salah satu agen di daftar.\n2. Periksa keberadaan dan navigasi antara tab Basic Info, Knowledge, dan Communication Style.',
        'expected_desc': '1. Form konfigurasi terbuka dengan 3 tab utama: Basic Info, Knowledge, dan Communication Style.\n2. User dapat berpindah antar tab dan seluruh form input ter-load tanpa error.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 3: Quick Setup AI Agent (US 3)'
    },
    {
        'unique_id': 'TC_AI_002',
        'title': 'Verifikasi penugasan nomor WhatsApp tunggal (Single WABA Assignment) pada AI Agent',
        'decription': 'Memastikan AI Agent hanya dapat ditugaskan ke 1 nomor WhatsApp (1-to-1) dan tidak dapat menambah multi-nomor, namun dapat mengganti nomor yang telah di-assign.',
        'precondition': '1. Buka konfigurasi AI Agent A.\n2. Terdapat lebih dari 1 nomor WABA terdaftar pada akun (misal: "123" dan "456").',
        'priority': 'P2 - High',
        'step_desc': '1. Klik "Pilih nomor Whatsapp" dan pilih nomor "123". Simpan.\n2. Buka kembali form penugasan nomor dan coba tambahkan nomor kedua "456".\n3. Ganti penugasan dari nomor "123" menjadi nomor "456". Simpan.',
        'expected_desc': '1. Sistem menetapkan nomor "123" pada AI Agent A.\n2. UI tidak mengizinkan penambahan multi-nomor (hanya single select dropdown).\n3. Penggantian ke nomor "456" berhasil dan nomor "123" terlepas dari AI Agent A.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 3: Quick Setup AI Agent (US 3)'
    },
    {
        'unique_id': 'TC_AI_003',
        'title': 'Verifikasi pesan peringatan prioritas alur Chat Flow saat mengaktifkan AI Autoreply',
        'decription': 'Memastikan sistem menampilkan warning message bahwa chat yang cocok dengan kondisi Chat Flow akan memprioritaskan alur Chat Flow saat toggle AI Autoreply diaktifkan.',
        'precondition': '1. Terdapat Chat Flow X pada nomor "123" dengan trigger Source A dan B.\n2. AI Agent A di-assign ke nomor "123".',
        'priority': 'P2 - High',
        'step_desc': '1. Buka konfigurasi AI Agent A.\n2. Aktifkan toggle "AI Auto-reply (Agentic)".\n3. Periksa notifikasi / warning modal yang muncul.',
        'expected_desc': '1. Tampil pesan peringatan: "New chat with matching conditions will follow the first response defined in Chat flow/Automation. AI agent A can still serve the chat if defined as a transfer destination.".\n2. Chat masuk dengan Source A/B dieksekusi oleh Chat Flow X; chat masuk dengan Source C dieksekusi oleh AI Agent A.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 3: Quick Setup AI Agent (US 3)'
    },

    # --- MODUL 4: Transfer Conditions & Fallback Settings (US 4) ---
    {
        'unique_id': 'TC_AI_004',
        'title': 'Verifikasi kewajiban pengisian Transfer Condition saat toggle AI Autoreply aktif',
        'decription': 'Memastikan bagian Transfer Condition wajib diisi (definisi intent & aksi handover ke Human Agent/Divisi) ketika toggle AI Autoreply aktif.',
        'precondition': '1. Buka form konfigurasi AI Agent A.\n2. Toggle "AI Auto-reply (Agentic)" diaktifkan ke ON.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Periksa bagian Transfer Condition yang muncul.\n2. Kosongkan intent / tujuan handover dan coba klik Simpan.\n3. Isi intent transfer dan pilih divisi/agen tujuan handover, lalu klik Simpan.',
        'expected_desc': '1. Bagian Transfer Condition muncul secara mandatory saat toggle ON.\n2. Form menampilkan validasi error wajib isi jika intent/destinasi kosong.\n3. Form berhasil disimpan saat intent dan destinasi handover terisi valid.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 4: Transfer & Fallback Condition (US 4)'
    },
    {
        'unique_id': 'TC_AI_005',
        'title': 'Verifikasi kelengkapan konfigurasi Fallback Condition pada AI Agent',
        'decription': 'Memastikan user diwajibkan menentukan destinasi handover ketika terjadi 4 kondisi fallback: Undefined Intent, AI Service Disruption / Timeout, AI Credit Exhausted, dan AI Repetitive Prevention.',
        'precondition': '1. Buka form konfigurasi AI Agent A.\n2. Toggle "AI Auto-reply (Agentic)" aktif (ON).',
        'priority': 'P1 - Critical',
        'step_desc': '1. Periksa bagian Fallback Condition pada form.\n2. Verifikasi 4 kondisi pemicu fallback: Undefined Intent, AI Disrupted/Timeout, AI Credit Exhausted, dan AI Repetitive (2x pesan identik berturut-turut).\n3. Tentukan destinasi handover untuk fallback dan simpan.',
        'expected_desc': '1. Bagian Fallback menampilkan 4 kondisi pemicu secara jelas.\n2. Destinasi handover fallback wajib dipilih (Divisi Tim CS / Personal Agent).\n3. Konfigurasi fallback berhasil tersimpan.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 4: Transfer & Fallback Condition (US 4)'
    },
    {
        'unique_id': 'TC_AI_006',
        'title': 'Verifikasi eksekusi Fallback Handover saat AI Repetitive Prevention terpicu (2x respon sama)',
        'decription': 'Memastikan percakapan otomatis di-handover ke Human Agent ketika AI memberikan respon identik 2 kali berturut-turut.',
        'precondition': '1. AI Agent A aktif menangani percakapan customer.\n2. Fallback dikonfigurasi ke Divisi CS Utama.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Kirim pesan customer yang memicu respon AI yang sama sebanyak 2 kali berturut-turut.\n2. Periksa status percakapan dan routing penanganan.',
        'expected_desc': '1. Sistem mendeteksi AI repetitive (2x respon identik).\n2. AI Agent berhenti merespon dan percakapan seketika di-handover ke antrean Divisi CS Utama sesuai aturan fallback.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 4: Transfer & Fallback Condition (US 4)'
    },
    {
        'unique_id': 'TC_AI_007',
        'title': 'Verifikasi penanganan AI Timeout / Disruption (Ambang batas 2 menit) pada runtime',
        'decription': 'Memastikan jika AI service tidak memberikan balasan selama 2 menit (timeout), sistem mencatat kondisi timeout dan memicu fallback handover ke CS.',
        'precondition': '1. Simulasi kondisi layanan AI mengalami gangguan / delay response > 2 menit.\n2. Percakapan customer masuk ke AI Agent.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Kirim pesan customer ke nomor WABA yang ditangani AI.\n2. Biarkan sistem menunggu hingga melewati ambang batas 2 menit tanpa balasan AI.\n3. Periksa status room chat.',
        'expected_desc': '1. Sistem mengidentifikasi AI Timeout (> 2 menit tanpa balasan).\n2. Percakapan dialihkan ke fallback Human Agent agar pesan customer tidak terabaikan.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 4: Transfer & Fallback Condition (US 4)'
    },

    # --- MODUL 5: AI Agent & Chat Flow Cross-alignment & Integrity (US 5) ---
    {
        'unique_id': 'TC_ALIGN_001',
        'title': 'Verifikasi pesan peringatan saat menonaktifkan AI Agent yang terhubung ke Chat Flow',
        'decription': 'Memastikan muncul dialog konfirmasi peringatan ketika user mencoba mematikan toggle AI Autoreply atau mencopot penugasan nomor WABA dari AI Agent yang sedang dipakai dalam Chat Flow.',
        'precondition': '1. AI Agent A terdaftar dan terlibat dalam Chat Flow X (aktif/inaktif) pada nomor "123".',
        'priority': 'P2 - High',
        'step_desc': '1. Buka konfigurasi AI Agent A.\n2. Matikan toggle "AI Auto-reply (Agentic)" (atau coba unassign nomor WABA "123").\n3. Periksa dialog peringatan yang muncul.\n4. Klik "Yes" pada dialog konfirmasi.',
        'expected_desc': '1. Muncul modal peringatan: "this AI Agent A is involved in chat flow/automation X, continue?".\n2. Saat user klik "Yes", status AI berhasil dimatikan/unassigned, dan Chat Flow X tetap dipertahankan di daftar dalam kondisi terhubung dengan AI yang inaktif.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 5: Cross-alignment & Integrity (US 5)'
    },
    {
        'unique_id': 'TC_ALIGN_002',
        'title': 'Verifikasi pencegahan re-aktivasi broken Chat Flow akibat AI Agent inaktif / unassigned',
        'decription': 'Memastikan sistem memeriksa integritas alur dan mencegah aktivasi Chat Flow jika AI Agent di dalamnya inaktif atau nomor WABA-nya telah dicopot.',
        'precondition': '1. Chat Flow X berstatus "Deactive" dan di dalamnya melibatkan AI Agent A.\n2. AI Agent A dalam kondisi inaktif atau unassigned dari nomor WABA.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Buka menu Pengaturan Alur Chat.\n2. Coba geser toggle status Chat Flow X menjadi Active.',
        'expected_desc': '1. Sistem melakukan pengecekan integritas alur (integrity check).\n2. Sistem memblokir aktivasi alur dan menampilkan pesan error spesifik mengenai alasan kegagalan (misal: "Alur tidak dapat diaktifkan karena AI Agent yang terhubung sedang tidak aktif").',
        'is_automated': 'FALSE',
        'user_story': 'Modul 5: Cross-alignment & Integrity (US 5)'
    },

    # --- MODUL 6: Human Agent Menu Restructuring (US 6) ---
    {
        'unique_id': 'TC_CS_001',
        'title': 'Verifikasi struktur sub-menu baru di bawah grup menu Human Agent / Tim CS',
        'decription': 'Memastikan menu Human Agent / Tim CS memuat 4 sub-menu terstruktur: Daftar Tim CS, Divisi, Jam kerja, dan Pengaturan distribusi pesan.',
        'precondition': '1. User login sebagai Owner/Supervisor.\n2. Buka menu navigasi Tim CS.',
        'priority': 'P3 - Medium',
        'step_desc': '1. Klik menu "Human Agent / Tim CS" di sidebar navigasi.\n2. Periksa daftar sub-menu yang tampil di dalamnya.',
        'expected_desc': '1. Menampilkan sub-menu secara berurutan:\n- Daftar Tim CS\n- Divisi\n- Jam kerja\n- Pengaturan distribusi pesan.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 6: Human Agent Menu Restructuring (US 6)'
    },
    {
        'unique_id': 'TC_CS_002',
        'title': 'Verifikasi rename menu "Pengaturan Chat Masuk" menjadi "Pengaturan Distribusi Pesan" & persistensi data',
        'decription': 'Memastikan menu telah di-rename dan seluruh konfigurasi distribusi eksisting (sticky agent, idle duration, rasio proporsi distribusi) tetap terjaga utuh.',
        'precondition': '1. Akun memiliki konfigurasi distribusi eksisting (sticky agent ON, proporsi agen terisi).',
        'priority': 'P2 - High',
        'step_desc': '1. Buka menu Human Agent -> Pengaturan Distribusi Pesan.\n2. Periksa nama header halaman.\n3. Periksa nilai pengaturan sticky agent, durasi idle, dan proporsi distribusi.',
        'expected_desc': '1. Nama menu dan judul halaman tertulis "Pengaturan Distribusi Pesan".\n2. Seluruh data konfigurasi distribusi eksisting user tetap utuh tanpa perubahan nilai.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 6: Human Agent Menu Restructuring (US 6)'
    },

    # --- MODUL 7: Human Agent & Chat Flow Distribution Alignment (US 7) ---
    {
        'unique_id': 'TC_CS_003',
        'title': 'Verifikasi pengalihan konfigurasi distribusi keyword ke menu Chat Flow',
        'decription': 'Memastikan konfigurasi distribusi berbasis kata kunci dialihkan dari menu Tim CS ke form Chat Flow dengan nomor WABA yang ter-prefill.',
        'precondition': '1. Buka menu Tim CS -> Pengaturan Distribusi Pesan untuk nomor "123".',
        'priority': 'P2 - High',
        'step_desc': '1. Klik opsi untuk mengatur distribusi ke divisi berdasarkan kata kunci.\n2. Periksa redirect halaman dan form yang terbuka.\n3. Konfigurasi keyword & response, lalu klik Simpan.',
        'expected_desc': '1. Pengguna diarahkan ke form pembuatan Chat Flow.\n2. Nomor WABA "123" terisi otomatis (prefilled) dan terkunci (read-only).\n3. Setelah disimpan, alur baru terbentuk dengan Source = default dan Keyword = custom.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 7: Human Agent & Chat Flow Alignment (US 7)'
    },
    {
        'unique_id': 'TC_CS_004',
        'title': 'Verifikasi migrasi otomatis data distribusi keyword eksisting menjadi entitas Chat Flow',
        'decription': 'Memastikan pasangan keyword-divisi eksisting (misal: 5 pasang) otomatis terkonversi menjadi entitas alur "CS Custom distribution [n]" di daftar Chat Flow.',
        'precondition': '1. Akun memiliki 5 konfigurasi distribusi kata kunci kustom lama sebelum pembaruan rilis.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Buka menu Pengaturan Alur Chat setelah fitur rilis.\n2. Periksa daftar alur yang terdaftar.',
        'expected_desc': '1. Ditemukan 5 alur chat baru bernama "CS Custom distribution 1" s/d "CS Custom distribution 5".\n2. Masing-masing alur memiliki Source = default, Keyword = custom, dengan pemetaan keyword group dan target divisi yang sesuai persis dengan data lama.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 7: Human Agent & Chat Flow Alignment (US 7)'
    },
    {
        'unique_id': 'TC_CS_005',
        'title': 'Verifikasi indikator status distribusi (Default vs Custom) pada menu Distribusi Pesan',
        'decription': 'Memastikan menu distribusi pesan menampilkan status "distribution custom" jika ada alur kustom di nomor tersebut, atau "distribution default" jika hanya ada flow primary default.',
        'precondition': '1. Terdapat nomor "123" dengan alur custom dan nomor "456" yang hanya memiliki primary default flow.',
        'priority': 'P3 - Medium',
        'step_desc': '1. Buka menu Pengaturan Distribusi Pesan untuk nomor "123". Periksa status tampilan.\n2. Buka menu Pengaturan Distribusi Pesan untuk nomor "456". Periksa status tampilan.',
        'expected_desc': '1. Pada nomor "123", sistem menampilkan status "distribution custom" aktif.\n2. Pada nomor "456", sistem menampilkan status "distribution default" aktif.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 7: Human Agent & Chat Flow Alignment (US 7)'
    },

    # --- MODUL 8: Visual Builder Handover Template to AI Agent (US 8) ---
    {
        'unique_id': 'TC_VB_001',
        'title': 'Verifikasi opsi handover dari kartu Template ke AI Agent pada Visual Builder',
        'decription': 'Memastikan kartu Template pada visual builder menyediakan opsi handover percakapan ke New AI Agent card maupun Existing AI Agent card.',
        'precondition': '1. Buka visual builder pada salah satu Chat Flow.\n2. Pilih sebuah kartu Template pada canvas.',
        'priority': 'P2 - High',
        'step_desc': '1. Klik titik koneksi / aksi handover pada kartu Template.\n2. Periksa opsi pilihan destinasi handover yang tersedia.',
        'expected_desc': '1. Muncul pilihan destinasi handover ke "AI Agent".\n2. Tersedia 2 sub-opsi: "New AI Agent card" dan "Existing AI Agent card".',
        'is_automated': 'FALSE',
        'user_story': 'Modul 8: Visual Builder Handover to AI Agent (US 8)'
    },
    {
        'unique_id': 'TC_VB_002',
        'title': 'Verifikasi pembuatan New AI Agent card secara inline dari dalam Visual Builder',
        'decription': 'Memastikan user dapat membuat AI Agent baru langsung dari visual builder dengan nomor WABA dan toggle terkunci.',
        'precondition': '1. Buka visual builder Chat Flow pada nomor "123".\n2. Pilih kartu Template dan klik Handover -> New AI Agent card.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Buka form pembuatan New AI Agent di dalam visual builder.\n2. Konfigurasi Name, Role, Goals, Description, Transfer Condition, dan Fallback.\n3. Periksa status field Phone Number dan toggle AI Auto-reply.\n4. Klik Simpan.',
        'expected_desc': '1. Form pembuatan AI Agent terbuka di dalam visual builder.\n2. Field nomor telepon terkunci pada "123" dan toggle AI Auto-reply terkunci aktif.\n3. Setelah disimpan, kartu AI Agent baru otomatis terhubung setelah kartu Template dan tersimpan dalam alur.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 8: Visual Builder Handover to AI Agent (US 8)'
    },
    {
        'unique_id': 'TC_VB_003',
        'title': 'Verifikasi koneksi handover kartu Template ke Existing AI Agent card tanpa duplikasi',
        'decription': 'Memastikan memilih Existing AI Agent menghubungkan kartu template ke kartu AI yang sudah ada tanpa menciptakan duplikasi kartu.',
        'precondition': '1. Terdapat kartu AI Agent A yang sudah ada di visual builder.\n2. Pilih kartu Template.',
        'priority': 'P2 - High',
        'step_desc': '1. Klik Handover pada kartu Template -> Pilih "Existing AI Agent card".\n2. Pilih "AI Agent A" dari daftar.\n3. Simpan konfigurasi flow.',
        'expected_desc': '1. Kartu Template berhasil terhubung ke kartu AI Agent A yang sudah ada.\n2. Sistem tidak membuat kartu AI Agent duplikat pada canvas visual builder.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 8: Visual Builder Handover to AI Agent (US 8)'
    },
    {
        'unique_id': 'TC_VB_004',
        'title': 'Verifikasi eksekusi runtime handover dari Template ke AI Agent dan transfer konteks chat',
        'decription': 'Memastikan saat percakapan customer mencapai titik handover kartu template, percakapan berpindah ke AI Agent dengan membawa konteks pesan sebelumnya.',
        'precondition': '1. Chat Flow aktif dengan susunan: Kartu Template -> Handover ke AI Agent A.\n2. Customer mengirim pesan masuk.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Customer menerima balasan dari Template Automasi.\n2. Customer merespon hingga mencapai titik handover.\n3. Periksa respon selanjutnya dan status handler di Chat Room.',
        'expected_desc': '1. Template berhenti merespon.\n2. AI Agent A melanjutkan percakapan dengan memahami konteks pesan sebelumnya.\n3. Status active handler di Chat Room diperbarui menjadi AI Agent A.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 8: Visual Builder Handover to AI Agent (US 8)'
    },
    {
        'unique_id': 'TC_VB_005',
        'title': 'Verifikasi dialog aktivasi paksa saat menghubungkan handover ke Inactive AI Agent',
        'decription': 'Memastikan sistem memberikan dialog konfirmasi untuk mengaktifkan AI Agent inaktif saat dipilih sebagai destinasi handover pada visual builder.',
        'precondition': '1. AI Agent B dalam kondisi inaktif.\n2. User menghubungkan kartu Template ke existing AI Agent B.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Pilih existing "AI Agent B" (inaktif) sebagai tujuan handover.\n2. Periksa pesan konfirmasi yang muncul.\n3. Klik "Yes" pada dialog konfirmasi dan simpan alur.',
        'expected_desc': '1. Muncul dialog warning: "this action will activate AI Agent B, continue?".\n2. Saat user klik "Yes", status AI Agent B dipaksa menjadi "Active".\n3. Alur Chat Flow berhasil disimpan tanpa error.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 8: Visual Builder Handover to AI Agent (US 8)'
    },

    # --- MODUL 9: Distinct Chat Room Message Bubbles & Sender Identification (US 9) ---
    {
        'unique_id': 'TC_ROOM_001',
        'title': 'Verifikasi gaya visual pembeda bubble chat untuk pesan Template di Chat Room',
        'decription': 'Memastikan pesan yang dikirim oleh Template Automasi menggunakan bubble style khusus Template dan memiliki penanda label yang jelas.',
        'precondition': '1. Buka Chat Room pada percakapan yang berisi pesan terkirim dari Template automasi.',
        'priority': 'P2 - High',
        'step_desc': '1. Buka room chat terkait.\n2. Amati bubble chat dari pesan template.',
        'expected_desc': '1. Pesan template menggunakan styling bubble khusus Template.\n2. Tampil label/indikator bahwa pesan dikirim oleh "Template".\n3. Gaya visual berbeda secara jelas dari bubble AI Agent dan Human Agent.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 9: Chat Room Bubbles & Sender ID (US 9)'
    },
    {
        'unique_id': 'TC_ROOM_002',
        'title': 'Verifikasi gaya visual pembeda bubble chat dan nama pengirim untuk pesan AI Agent',
        'decription': 'Memastikan pesan yang dikirim oleh AI Agent menggunakan bubble style AI dan menampilkan nama AI Agent yang bersangkutan.',
        'precondition': '1. Buka Chat Room pada percakapan yang berisi pesan balasan dari AI Agent A.',
        'priority': 'P2 - High',
        'step_desc': '1. Buka room chat terkait.\n2. Amati bubble chat dari pesan AI Agent.',
        'expected_desc': '1. Pesan menggunakan styling bubble khusus AI Agent.\n2. Tampil label pengirim "AI Agent" beserta nama spesifik agen (misal: "AI Agent A").\n3. Tampilan visually distinct dari bubble Template dan Human CS.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 9: Chat Room Bubbles & Sender ID (US 9)'
    },
    {
        'unique_id': 'TC_ROOM_003',
        'title': 'Verifikasi gaya visual pembeda bubble chat dan nama agen untuk pesan Human Agent (CS)',
        'decription': 'Memastikan pesan yang dikirim manual oleh CS menggunakan bubble style Human Agent dan menampilkan nama agen CS pengirim.',
        'precondition': '1. Buka Chat Room pada percakapan yang dibalas manual oleh agen CS "Budi".',
        'priority': 'P2 - High',
        'step_desc': '1. Buka room chat terkait.\n2. Amati bubble chat dari balasan manual agen CS.',
        'expected_desc': '1. Pesan menggunakan styling bubble khusus Human Agent.\n2. Tampil label pengirim beserta nama agen CS "Budi".\n3. Tampilan visually distinct dari bubble Template dan AI Agent.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 9: Chat Room Bubbles & Sender ID (US 9)'
    },
    {
        'unique_id': 'TC_ROOM_004',
        'title': 'Verifikasi konsistensi bubble chat Customer dan persistensi riwayat chat (History)',
        'decription': 'Memastikan bubble chat customer tetap standar dan seluruh tipe bubble pada riwayat chat tidak berubah meskipun percakapan telah ditransfer ke handler lain.',
        'precondition': '1. Percakapan telah melalui fase: Customer chat -> Template balas -> AI balas -> CS balas manual -> Sesi ditutup.',
        'priority': 'P2 - High',
        'step_desc': '1. Buka kembali riwayat percakapan yang telah selesai tersebut.\n2. Amati seluruh bubble pesan dari awal hingga akhir percakapan.',
        'expected_desc': '1. Bubble chat customer tetap pada styling standar customer.\n2. Seluruh tipe bubble (Template, AI, Human CS) tetap mempertahankan gaya visual dan identitas pengirim aslinya (tipe sender tidak berubah/tertimpa).',
        'is_automated': 'FALSE',
        'user_story': 'Modul 9: Chat Room Bubbles & Sender ID (US 9)'
    },

    # --- MODUL 10: Chat Room Takeover & Turn Off Flow (US 10) ---
    {
        'unique_id': 'TC_TO_001',
        'title': 'Verifikasi tampilan tombol overlay Ambil Alih Chat pada percakapan yang ditangani AI/Flow',
        'decription': 'Memastikan tombol overlay "Ambil Alih Chat" tampil pada room percakapan yang sedang aktif ditangani oleh AI Agent atau Chat Flow.',
        'precondition': '1. Login sebagai Owner / Supervisor / CS Agent yang memiliki izin.\n2. Buka percakapan yang sedang berjalan di bawah kendali AI Agent / Chat Flow.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Buka room percakapan aktif di Chat Room.\n2. Periksa keberadaan tombol aksi di bagian bawah / overlay chat room.',
        'expected_desc': '1. Tombol aksi "Ambil Alih Chat" (Turn off flow) tampil jelas di atas area input percakapan bagi user yang berwenang.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 10: Chat Room Takeover & Turn Off Flow (US 10)'
    },
    {
        'unique_id': 'TC_TO_002',
        'title': 'Verifikasi eksekusi aksi Ambil Alih Chat (Penghentian respon otomatis & pembukaan input CS)',
        'decription': 'Memastikan saat tombol Ambil Alih Chat diklik, respon otomatis AI/flow seketika berhenti, tombol overlay hilang, dan input box terbuka untuk penanganan CS.',
        'precondition': '1. Buka room percakapan yang sedang ditangani AI Agent.\n2. Tombol "Ambil Alih Chat" tampil aktif.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Klik tombol "Ambil Alih Chat".\n2. Periksa perubahan status room dan fungsionalitas input box.\n3. Ketik pesan balasan manual dan kirim ke customer.',
        'expected_desc': '1. Active Chat Flow pada percakapan tersebut seketika berstatus OFF (turned off).\n2. Respon otomatis AI / langkah flow berikutnya berhenti merespon.\n3. Tombol overlay "Ambil Alih Chat" hilang.\n4. Input box terbuka aktif dan agen CS berhasil mengirim balasan manual ke customer.\n5. Status active handler di Chat Room diperbarui menjadi Human CS.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 10: Chat Room Takeover & Turn Off Flow (US 10)'
    },
    {
        'unique_id': 'TC_TO_003',
        'title': 'Verifikasi pencegahan trigger automasi ulang pada percakapan yang telah diambil alih (Single Session Guard)',
        'decription': 'Memastikan pesan baru dari customer dalam sesi yang sama setelah di-takeover tidak memicu flow automasi/AI kembali.',
        'precondition': '1. Percakapan telah diambil alih oleh CS (flow turned off).\n2. Sesi percakapan masih dalam status aktif/open.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Kirim pesan baru dari WhatsApp customer dengan keyword/intent yang biasanya memicu automasi atau AI.\n2. Periksa apakah sistem merespon secara otomatis atau tetap dalam kendali CS.',
        'expected_desc': '1. Sistem tidak memicu automasi atau AI Agent kembali pada percakapan tersebut.\n2. Pesan customer masuk ke antrean CS dan percakapan tetap berada di bawah penanganan manual Human Agent.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 10: Chat Room Takeover & Turn Off Flow (US 10)'
    },
    {
        'unique_id': 'TC_TO_004',
        'title': 'Verifikasi inisiasi ulang Chat Flow pada sesi percakapan baru setelah chatroom closing/resolved',
        'decription': 'Memastikan setelah sesi chatroom di-closing/resolved oleh CS, pesan baru yang dikirim customer di kemudian hari akan memicu Chat Flow dari awal kembali.',
        'precondition': '1. Percakapan sebelumnya telah diambil alih oleh CS dan telah diselesaikan (Closed / Resolved).\n2. Nomor WABA memiliki Chat Flow aktif.',
        'priority': 'P1 - Critical',
        'step_desc': '1. Pastikan tiket/sesi room chat sebelumnya telah berstatus Closed.\n2. Customer mengirimkan pesan WhatsApp baru untuk memulai percakapan baru.\n3. Periksa perilaku penanganan pesan masuk pertama pada sesi baru tersebut.',
        'expected_desc': '1. Sistem mendeteksi pesan sebagai sesi percakapan baru.\n2. Chat Flow / AI Agent terpicu kembali dari awal sesuai konfigurasi trigger first responder.\n3. Pesan otomatis awal dikirimkan kembali kepada customer.',
        'is_automated': 'FALSE',
        'user_story': 'Modul 10: Chat Room Takeover & Turn Off Flow (US 10)'
    }
]

for row_idx, tc in enumerate(test_cases, start=2):
    ws.cell(row=row_idx, column=1, value=None)
    ws.cell(row=row_idx, column=2, value=None)
    ws.cell(row=row_idx, column=3, value=tc['unique_id'])
    ws.cell(row=row_idx, column=4, value=tc['title'])
    ws.cell(row=row_idx, column=5, value=tc['decription'])
    ws.cell(row=row_idx, column=6, value=tc['precondition'])
    ws.cell(row=row_idx, column=7, value=tc['priority'])
    ws.cell(row=row_idx, column=8, value=None)
    ws.cell(row=row_idx, column=9, value=None)
    ws.cell(row=row_idx, column=10, value=None)
    ws.cell(row=row_idx, column=11, value=None)
    ws.cell(row=row_idx, column=12, value=tc['step_desc'])
    ws.cell(row=row_idx, column=13, value=tc['expected_desc'])
    ws.cell(row=row_idx, column=14, value=tc['is_automated'])
    ws.cell(row=row_idx, column=15, value=tc['user_story'])

    for col_idx in range(1, 16):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.border = thin_border
        if col_idx in [4, 5, 6, 12, 13, 15]:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        elif col_idx in [1, 2, 3, 7, 8, 9, 10, 11, 14]:
            cell.alignment = Alignment(horizontal='center', vertical='top')

out_path = '/Users/eldofadlyady/Documents/eldoTest/result-testcase/Test_Cases_Chat_Flow_Orchestration.xlsx'
wb.save(out_path)
print(f'Successfully generated {len(test_cases)} test cases to {out_path}')
