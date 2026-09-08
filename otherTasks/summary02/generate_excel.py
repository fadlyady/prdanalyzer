import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_test_cases():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    headers = [
        "project_id",
        "suite_id",
        "unique_id",
        "title",
        "decription",
        "precondition",
        "priority",
        "type",
        "status",
        "tags",
        "steps",
        "step_desc",
        "expected_desc",
        "is_automated",
        "user_story"
    ]

    ws.append(headers)

    # Styling header
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    test_cases = [
        # ==========================================
        # TS01: Pre-requisites & Role Access Control (RBAC)
        # ==========================================
        {
            "title": "Verifikasi pencegahan akses fitur Everpro Chat ketika Subscription tidak aktif / kedaluwarsa",
            "decription": "Memastikan akun merchant yang tidak memiliki paket subscription aktif dicegah untuk mengakses fitur Everpro Chat, WABA, dan pengiriman pesan berbayar.",
            "precondition": "1. Akun pengguna terdaftar di Everpro.\n2. Status paket Subscription akun adalah 'Inactive' atau 'Expired'.",
            "priority": "P1 - Critical",
            "step_desc": "1. Login ke Everpro menggunakan akun dengan subscription inactive.\n2. Navigasikan ke menu Everpro Chat / WhatsApp Credit / Broadcast.\n3. Coba lakukan inisiasi pengiriman pesan atau broadcast.",
            "expected_desc": "1. Sistem menampilkan banner/modal bahwa paket subscription tidak aktif.\n2. Pengguna diblokir untuk mengakses menu Everpro Chat dan tidak dapat mengirim pesan berbayar.\n3. Muncul CTA untuk mengaktifkan/memperpanjang subscription.",
            "user_story": "Global Pre-condition"
        },
        {
            "title": "Verifikasi kegagalan pengiriman pesan saat nomor WABA belum terdaftar atau berstatus non-aktif",
            "decription": "Memastikan sistem memvalidasi status WABA sebelum pengiriman pesan dilakukan (tidak ada proses isolasi/pengiriman jika WABA tidak valid).",
            "precondition": "1. Subscription akun berstatus Aktif.\n2. Organisasi belum mendaftarkan nomor WABA atau status WABA dalam peninjauan/banned.",
            "priority": "P1 - Critical",
            "step_desc": "1. Login sebagai Owner/Supervisor.\n2. Akses halaman Broadcast atau Chatroom.\n3. Coba kirim pesan template atau pesan chatroom ke kontak customer.",
            "expected_desc": "1. Sistem menampilkan peringatan bahwa WABA belum terdaftar/aktif.\n2. Pengiriman pesan dicegah dan tidak ada pemotongan/isolasi saldo WhatsApp Credit.",
            "user_story": "Global Pre-condition"
        },
        {
            "title": "Verifikasi kegagalan pengiriman pesan pada WABA yang belum memiliki agent yang di-assign",
            "decription": "Memastikan pesan tidak dapat diproses jika pada nomor WABA aktif belum ditugaskan (assigned) agent operasional.",
            "precondition": "1. Subscription akun Aktif.\n2. WABA terdaftar dan aktif.\n3. Belum ada CS/Agent yang di-assign ke nomor WABA tersebut.",
            "priority": "P1 - Critical",
            "step_desc": "1. Login sebagai Owner/Supervisor.\n2. Buka menu Chatroom untuk nomor WABA tanpa agent tersebut.\n3. Coba lakukan inisiasi percakapan/kirim pesan.",
            "expected_desc": "1. Sistem menampilkan validasi error bahwa agent belum ditugaskan ke WABA.\n2. Pengiriman pesan tidak dapat dieksekusi sampai minimal 1 agent di-assign.",
            "user_story": "Global Pre-condition"
        },
        {
            "title": "Verifikasi pembatasan hak akses role CS Operasional / CS Closing terhadap menu Admin Biz Ops & Credit Management",
            "decription": "Memastikan role CS tidak memiliki akses ke halaman Biz Ops Pricing, konfigurasi WhatsApp Credit, Monthly Credit Limit, dan Broadcast Management.",
            "precondition": "1. Subscription Aktif, WABA Aktif dengan agent assigned.\n2. Akun pengguna memiliki role 'CS Operasional' atau 'CS Closing'.",
            "priority": "P1 - Critical",
            "step_desc": "1. Login sebagai CS Operasional / CS Closing.\n2. Periksa menu navigasi yang tersedia.\n3. Coba akses direct URL halaman Admin Pricing (`/admin/pricing`) atau Credit Management (`/chat/credit-management`).",
            "expected_desc": "1. Menu Admin Pricing dan Credit Management tidak muncul pada navigasi CS.\n2. Direct URL akses menghasilkan halaman 403 Forbidden atau diarahkan kembali ke Chatroom.\n3. CS tidak dapat mengubah limit bulanan maupun harga.",
            "user_story": "Global RBAC Rule"
        },

        # ==========================================
        # TS02: Admin Pricing & Effective Date Engine (User Story 1)
        # ==========================================
        {
            "title": "Biz Ops Admin memperbarui Meta Base Price per Kategori x Negara Penerima dengan Effective Date",
            "decription": "Memastikan Biz Ops Admin dapat memperbarui base price resmi dari Meta untuk setiap kategori pesan (Marketing, Utility, Auth, Service) per negara tujuan beserta tanggal efektif berlakunya.",
            "precondition": "1. Login sebagai Everpro Biz Ops Admin.\n2. Memiliki otorisasi manajemen harga di database/admin dashboard.",
            "priority": "P2 - High",
            "step_desc": "1. Buka dashboard Credit Management Admin > Base Price.\n2. Pilih Kategori Pesan: 'Marketing', Negara Tujuan: 'Indonesia (ID)'.\n3. Masukkan harga dasar baru (misal: Rp450).\n4. Set Effective Date (misal: H+1 jam 00:00 WIB).\n5. Klik 'Simpan'.",
            "expected_desc": "1. Base price baru berhasil tersimpan di sistem.\n2. Sistem mencatat status pending aktif hingga tanggal efektif tercapai.\n3. Sebelum tanggal efektif, harga lama tetap digunakan.",
            "user_story": "User Story 1 - AC1"
        },
        {
            "title": "Biz Ops Admin mengatur General Markup Price per Kategori x Negara Penerima untuk seluruh user umum",
            "decription": "Memastikan Biz Ops Admin dapat menentukan margin/markup per pesan untuk seluruh general users berdasarkan kombinasi kategori pesan dan negara tujuan.",
            "precondition": "1. Login sebagai Everpro Biz Ops Admin.\n2. Base price untuk kategori & negara terkait sudah terkonfigurasi.",
            "priority": "P2 - High",
            "step_desc": "1. Buka menu Admin > Markup Price General.\n2. Pilih Kategori: 'Service', Negara: 'Indonesia (ID)'.\n3. Input nilai markup (misal: Rp50).\n4. Set Effective Date.\n5. Klik 'Simpan'.",
            "expected_desc": "1. Markup price general tersimpan dengan benar.\n2. Total tarif untuk general user terhitung sebagai (Base Price + Rp50) saat tanggal efektif aktif.",
            "user_story": "User Story 1 - AC2"
        },
        {
            "title": "Biz Ops Admin mengatur Customizable Markup Price untuk User ID spesifik (Special Merchant)",
            "decription": "Memastikan Biz Ops Admin dapat memberikan harga markup khusus (custom pricing) untuk merchant tertentu berdasarkan User ID.",
            "precondition": "1. Login sebagai Everpro Biz Ops Admin.\n2. Target User ID merchant valid dan terdaftar di database.",
            "priority": "P2 - High",
            "step_desc": "1. Buka menu Admin > Custom User Pricing.\n2. Masukkan target User ID (misal: `USR_100234`).\n3. Pilih Kategori: 'Marketing', Negara: 'Indonesia'.\n4. Masukkan Custom Markup (misal: Rp20, lebih rendah dari General Markup Rp50).\n5. Set Effective Date dan simpan.",
            "expected_desc": "1. Custom pricing untuk User ID tersebut berhasil tersimpan.\n2. Saat User ID `USR_100234` mengirim pesan, sistem menggunakan markup Rp20 (bukan markup general Rp50).",
            "user_story": "User Story 1 - AC3"
        },
        {
            "title": "Verifikasi masa transisi harga sebelum vs sesudah Effective Date tercapai",
            "decription": "Memastikan sistem menerapkan harga lama sebelum Effective Date dan beralih ke harga baru secara otomatis tepat saat Effective Date tiba.",
            "precondition": "1. Konfigurasi harga baru dengan Effective Date: 2026-10-01 00:00:00 WIB.\n2. Harga lama: Rp500, Harga baru: Rp600.",
            "priority": "P2 - High",
            "step_desc": "1. Kirim pesan berbayar pada 2026-09-30 23:59:50 WIB.\n2. Cek nominal kredit yang diisolasi/dipotong.\n3. Kirim pesan berbayar pada 2026-10-01 00:00:05 WIB.\n4. Cek nominal kredit yang diisolasi/dipotong.",
            "expected_desc": "1. Pesan pada step 1 dikenakan tarif lama (Rp500).\n2. Pesan pada step 3 dikenakan tarif baru (Rp600).\n3. Transisi pergantian harga berjalan mulus tanpa downtime.",
            "user_story": "User Story 1 - AC1, AC2, AC3"
        },
        {
            "title": "Verifikasi kelengkapan pencatatan Audit Log saat terjadi perubahan harga oleh Biz Ops Admin",
            "decription": "Memastikan sistem mencatat audit log lengkap setiap kali terjadi penambahan/perubahan base price, general markup, atau custom user pricing.",
            "precondition": "1. Login sebagai Everpro Biz Ops Admin.\n2. Melakukan perubahan harga pada salah satu kategori.",
            "priority": "P3 - Medium",
            "step_desc": "1. Lakukan update harga markup untuk User ID `USR_8899` dari Rp30 menjadi Rp40.\n2. Buka menu Audit Log / Database Log Perubahan Harga.\n3. Periksa baris log terbaru.",
            "expected_desc": "1. Terbentuk baris log baru dengan informasi lengkap:\n   - Actor (ID/Email Admin yang mengubah)\n   - New price & Old price\n   - Message category & Recipient country\n   - Target User ID\n   - Timestamp perubahan dan Timestamp tanggal efektif.",
            "user_story": "User Story 1 - AC4"
        },

        # ==========================================
        # TS03: Free Entry Point (FEP) Window Validation (User Story 2)
        # ==========================================
        {
            "title": "Verifikasi pesan inbound dari CTWA (Click-to-WhatsApp Ads) mengaktifkan Free Entry Point (FEP) window",
            "decription": "Memastikan saat customer mengklik iklan CTWA dan mengirim pesan masuk, sistem menandai window percakapan tersebut sebagai FEP aktif (bebas biaya).",
            "precondition": "1. Nomor WABA terintegrasi dengan Meta Ads (CTWA).\n2. Saldo kredit WhatsApp merchant mencukupi.",
            "priority": "P1 - Critical",
            "step_desc": "1. Customer mengirimkan pesan masuk pertama kali via trigger CTWA.\n2. Periksa status session di sistem backend / chatroom.",
            "expected_desc": "1. Sistem mengenali payload referral CTWA dari Meta.\n2. Flag `Free Entry Point Window` aktif untuk kontak/nomor customer tersebut sesuai durasi Meta window.",
            "user_story": "User Story 2 - AC1"
        },
        {
            "title": "Verifikasi pesan inbound dari WTWA (Website-to-WhatsApp) mengaktifkan Free Entry Point (FEP) window",
            "decription": "Memastikan saat customer mengirim pesan melalui button WTWA pada website resmi, sistem mengaktifkan FEP window.",
            "precondition": "1. Nomor WABA terpasang button WTWA.\n2. Customer mengirim pesan masuk melalui link WTWA.",
            "priority": "P1 - Critical",
            "step_desc": "1. Kirim inbound pesan dari customer via WTWA.\n2. Verifikasi status session percakapan di sistem.",
            "expected_desc": "1. Sistem mendeteksi referral WTWA.\n2. FEP window aktif untuk sesi percakapan tersebut.",
            "user_story": "User Story 2 - AC1"
        },
        {
            "title": "Pengiriman pesan Template (Marketing/Utility/Auth/Service) via Broadcast selama window FEP aktif tanpa pemotongan kredit",
            "decription": "Memastikan jika target kontak penerima broadcast sedang berada dalam window FEP aktif, pesan langsung dikirim tanpa validasi saldo dan tanpa isolasi/pemotongan kredit.",
            "precondition": "1. Kontak customer memiliki status FEP window aktif.\n2. Merchant membuat campaign broadcast template Marketing.",
            "priority": "P1 - Critical",
            "step_desc": "1. Buat dan jadwalkan broadcast campaign ke kontak tersebut.\n2. Eksekusi pengiriman broadcast.\n3. Periksa saldo WhatsApp Credit merchant dan log transaksi.",
            "expected_desc": "1. Pesan broadcast berhasil dikirim ke Meta.\n2. Saldo WhatsApp Credit merchant TIDAK berkurang / tidak diisolasi (Rp0).\n3. Pesan terkirim segera tanpa terhambat validasi saldo.",
            "user_story": "User Story 2 - AC1"
        },
        {
            "title": "Pengiriman pesan balasan via Chatroom (Human / AI Bot) dan Open API selama window FEP aktif bebas biaya",
            "decription": "Memastikan agen CS, AI bot, atau trigger Open API yang mengirim pesan ke kontak dalam window FEP tidak dikenakan biaya kredit.",
            "precondition": "1. Percakapan chatroom customer dalam status FEP aktif.",
            "priority": "P1 - Critical",
            "step_desc": "1. CS mengirimkan pesan manual di chatroom.\n2. AI bot merespon pertanyaan customer.\n3. Sistem backend men-trigger pengiriman pesan via Open API.\n4. Cek saldo kredit dan running cost chatroom.",
            "expected_desc": "1. Seluruh pesan terkirim sukses ke customer.\n2. Tidak ada kredit yang diisolasi/dipotong.\n3. Running cost session di chatroom tetap Rp0.",
            "user_story": "User Story 2 - AC1"
        },
        {
            "title": "Verifikasi pengiriman pesan setelah FEP window kedaluwarsa (Expired) kembali dikenakan pemotongan kredit normal",
            "decription": "Memastikan setelah masa aktif window FEP dari Meta berakhir, pesan berikutnya yang dikirimkan kembali melewati proses isolasi dan pemotongan kredit.",
            "precondition": "1. Kontak customer sebelumnya memiliki FEP aktif.\n2. Masa aktif window FEP telah habis/expired.",
            "priority": "P1 - Critical",
            "step_desc": "1. Kirim pesan berbayar (Service Message / Template) ke customer setelah FEP window habis.\n2. Amati proses pada sistem dan saldo kredit.",
            "expected_desc": "1. Sistem mendeteksi FEP window sudah tidak aktif.\n2. Sistem menjalankan validasi saldo dan mengisolasi kredit sebesar `Base Price + Margin`.",
            "user_story": "User Story 2 - AC1, AC2"
        },

        # ==========================================
        # TS04: Credit Isolation & Broadcast Balance Validation (User Story 2)
        # ==========================================
        {
            "title": "Verifikasi kalkulasi akurat nilai isolasi kredit (Meta Base Price + Everpro Margin)",
            "decription": "Memastikan besaran nominal kredit yang diisolasi sementara sesuai dengan rumus: Base Price (sesuai template type & negara) + Margin Everpro (sesuai User ID/general markup).",
            "precondition": "1. User ID `USR_01` dengan custom margin Marketing ID = Rp25.\n2. Base price Marketing ID = Rp450.\n3. Total estimasi per pesan = Rp475.",
            "priority": "P1 - Critical",
            "step_desc": "1. Kirim 1 pesan template Marketing ke nomor Indonesia di luar FEP window.\n2. Periksa mutasi saldo kredit yang diisolasi di database/sistem.",
            "expected_desc": "1. Sistem mengisolasi saldo sebesar tepat Rp475.\n2. Saldo aktif yang tersedia berkurang Rp475 (status temporary hold/deduct).",
            "user_story": "User Story 2 - AC2"
        },
        {
            "title": "Verifikasi eksekusi Broadcast dengan saldo mencukupi berhasil mengisolasi kredit seluruh target kontak",
            "decription": "Memastikan broadcast campaign dengan saldo yang cukup berhasil mengisolasi total kredit sesuai `jumlah kontak non-FEP * tarif per pesan`.",
            "precondition": "1. Saldo WhatsApp Credit: Rp1.000.000.\n2. Broadcast ditujukan ke 1.000 kontak non-FEP dengan tarif Rp500/pesan (Total kebutuhan: Rp500.000).",
            "priority": "P1 - Critical",
            "step_desc": "1. Jadwalkan broadcast campaign.\n2. Tunggu hingga waktu jadwal broadcast tiba / klik kirim sekarang.\n3. Periksa saldo kredit dan status antrian pengiriman.",
            "expected_desc": "1. Sistem berhasil mengisolasi saldo sebesar Rp500.000.\n2. Saldo tersedia menjadi Rp500.000.\n3. Pengiriman 1.000 pesan ke Meta dimulai.",
            "user_story": "User Story 2 - AC2"
        },
        {
            "title": "Verifikasi validasi 'All-or-Nothing' pada Broadcast dengan saldo parsial / tidak mencukupi (Insufficient Credit)",
            "decription": "Memastikan sistem langsung MENOLAK broadcast secara utuh di awal jika saldo kredit tidak cukup menutupi total estimasi biaya seluruh kontak, tanpa melakukan pengiriman parsial.",
            "precondition": "1. Saldo WhatsApp Credit: Rp300.000.\n2. Broadcast dibuat untuk 1.000 kontak dengan tarif Rp500/pesan (Total kebutuhan: Rp500.000).",
            "priority": "P1 - Critical",
            "step_desc": "1. Simpan dan jalankan broadcast campaign.\n2. Amati respon sistem saat proses validasi saldo awal.",
            "expected_desc": "1. Sistem langsung menolak eksekusi broadcast di awal.\n2. Muncul pesan error 'Saldo WhatsApp Credit tidak mencukupi untuk mengirim campaign'.\n3. Tidak ada pesan yang terkirim ke Meta dan tidak ada saldo yang diisolasi (saldo tetap Rp300.000).",
            "user_story": "User Story 2 - AC2"
        },
        {
            "title": "Verifikasi isolasi kredit instan saat CS mengklik tombol balas di Chatroom atau pemicuan Open API",
            "decription": "Memastikan pengiriman pesan berbayar dari Chatroom atau Open API langsung mengisolasi kredit sesaat sebelum pesan diteruskan ke Meta.",
            "precondition": "1. Saldo kredit merchant mencukupi.\n2. Percakapan chatroom di luar FEP window.",
            "priority": "P1 - Critical",
            "step_desc": "1. CS mengetik balasan di chatroom dan menekan tombol kirim.\n2. Amati lifecycle kredit pada backend.",
            "expected_desc": "1. Sistem langsung mengisolasi nilai kredit untuk pesan tersebut.\n2. Pesan dikirimkan ke Meta gateway.",
            "user_story": "User Story 2 - AC2"
        },

        # ==========================================
        # TS05: Credit Rollback & Nett Credit Logging (User Story 2)
        # ==========================================
        {
            "title": "Verifikasi Credit Rollback otomatis saat pesan menerima sinyal eksplisit status 'failed' dari Meta",
            "decription": "Memastikan sistem mengembalikan (rollback) saldo kredit yang diisolasi sejumlah 100% dari nilai isolasi awal ketika Meta merespon status gagal.",
            "precondition": "1. Pesan telah diisolasi kreditnya sebesar Rp500 dan dikirim ke Meta.\n2. Nomor tujuan tidak valid atau Meta mengembalikan status webhook 'failed'.",
            "priority": "P1 - Critical",
            "step_desc": "1. Simulasikan webhook callback dari Meta dengan status pesan 'failed'.\n2. Periksa saldo kredit merchant dan log mutasi.",
            "expected_desc": "1. Sistem memicu proses credit rollback sebesar Rp500 secara instan.\n2. Saldo kredit merchant bertambah kembali Rp500.\n3. Status pesan tercatat 'failed' dengan alasan kegagalan dari Meta.",
            "user_story": "User Story 2 - AC2"
        },
        {
            "title": "Verifikasi perhitungan Nett Credit Log (Temporary Deducted - Rolled Back) pada status akhir",
            "decription": "Memastikan pencatatan nett credit log akhir akurat sesuai status akhir pesan (terkirim sukses vs gagal).",
            "precondition": "1. Terjadi pengiriman broadcast multi-kontak (sebagian sukses 'delivered', sebagian 'failed').",
            "priority": "P1 - Critical",
            "step_desc": "1. Kirim campaign ke 10 kontak (8 sukses delivered, 2 failed).\n2. Amati kalkulasi nett credit log di backend.",
            "expected_desc": "1. Untuk 8 pesan sukses: Nett deduction = Rp500/pesan (Total Rp4.000).\n2. Untuk 2 pesan gagal: Temporary deduction (Rp1.000) - Rollback (Rp1.000) = Nett deduction Rp0.\n3. Total saldo terpotong akhir adalah tepat Rp4.000.",
            "user_story": "User Story 2 - AC3"
        },
        {
            "title": "Verifikasi integritas data log transaksi kredit dapat dilihat dan diekspor pada fitur Credit History",
            "decription": "Memastikan seluruh mutasi kredit (isolasi, rollback, final deduction) tersimpan dengan rapi dan dapat diakses pada riwayat kredit.",
            "precondition": "1. Telah ada transaksi pengiriman pesan berbayar dan rollback.",
            "priority": "P2 - High",
            "step_desc": "1. Buka menu WhatsApp Credit > Credit History.\n2. Periksa daftar log transaksi yang tampil.",
            "expected_desc": "1. Daftar riwayat kredit menampilkan rincian transaksi dengan benar (ID Transaksi, WABA, Nominal, Status, Tanggal).\n2. Data konsisten antara tampilan UI dan database.",
            "user_story": "User Story 2 - AC3"
        },

        # ==========================================
        # TS06: Service Message Expense & Monthly Limit (User Story 3)
        # ==========================================
        {
            "title": "Verifikasi 'Pengeluaran bulan ini' pada WhatsApp Credit Page mengikutsertakan biaya Service Message",
            "decription": "Memastikan widget kalkulasi pengeluaran bulanan mencakup seluruh kategori pesan berbayar: Marketing, Utility, Authentication, dan Service Message.",
            "precondition": "1. Dalam bulan berjalan terdapat pengeluaran:\n   - Marketing: Rp100.000\n   - Utility: Rp50.000\n   - Auth: Rp20.000\n   - Service Message: Rp80.000",
            "priority": "P2 - High",
            "step_desc": "1. Login sebagai Owner/Supervisor.\n2. Buka halaman WhatsApp Credit.\n3. Periksa angka pada kartu 'Pengeluaran bulan ini' untuk nomor WABA tersebut.",
            "expected_desc": "1. Total 'Pengeluaran bulan ini' menampilkan akumulasi yang tepat: Rp250.000 (termasuk Service Message Rp80.000).\n2. Rincian per kategori menunjukkan porsi Service Message dengan jelas.",
            "user_story": "User Story 3 - AC1"
        },
        {
            "title": "Verifikasi pemblokiran pengiriman pesan saat akumulasi biaya mencapai Monthly Credit Limit",
            "decription": "Memastikan sistem mencegah pengiriman pesan berbayar apapun ketika kumulatif pengeluaran (volume * tarif) telah mencapai atau melebihi limit bulanan yang disetel.",
            "precondition": "1. Owner telah mengatur 'Monthly Credit Limit' = Rp1.000.000.\n2. Total pengeluaran bulan berjalan telah mencapai Rp1.000.000.",
            "priority": "P2 - High",
            "step_desc": "1. Coba kirim pesan template berbayar atau balasan service message baru.\n2. Amati respon sistem.",
            "expected_desc": "1. Sistem menolak pengiriman pesan.\n2. Muncul pesan sistem: 'monthly credit limit has exhausted'.\n3. Tidak ada penambahan biaya melebihi limit yang ditentukan.",
            "user_story": "User Story 3 - AC2"
        },
        {
            "title": "Verifikasi ekspor file Riwayat Kredit (Credit History Download) memuat baris Service Message",
            "decription": "Memastikan file hasil unduh riwayat pemakaian kredit menyertakan baris transaksi Service Message dengan nilai kolom TEMPLATE TYPE = 'Service'.",
            "precondition": "1. WABA memiliki riwayat pengiriman pesan Service, Marketing, Utility, dan Auth.",
            "priority": "P2 - High",
            "step_desc": "1. Buka menu WhatsApp Credit > Download Usage History.\n2. Pilih periode bulan terkait dan unduh file (CSV/Excel).\n3. Buka file hasil unduhan dan periksa baris transaksi.",
            "expected_desc": "1. File berhasil diunduh tanpa korupsi data.\n2. Transaksi Service Message tercantum dengan jelas.\n3. Nilai kolom 'TEMPLATE TYPE' tertulis 'Service', dapat dibedakan dari Marketing, Utility, dan Authentication.\n4. Tidak ada kolom tambahan di luar spesifikasi yang disetujui.",
            "user_story": "User Story 3 - AC3"
        },

        # ==========================================
        # TS07: Chatroom Session Cost & Flat Limit Setting (User Story 4)
        # ==========================================
        {
            "title": "Verifikasi running cost sesi aktif di Chatroom ter-update secara real-time saat pesan berbayar terkirim",
            "decription": "Memastikan UI chatroom menampilkan running cost sesi yang bertambah secara real-time setiap kali pesan berbayar terkirim dalam sesi 24 jam Meta aktif.",
            "precondition": "1. Sesi chatroom aktif dengan customer di luar FEP.\n2. Running cost awal = Rp0.",
            "priority": "P3 - Medium",
            "step_desc": "1. Agen mengirim pesan berbayar pertama (biaya Rp500).\n2. Amati tampilan widget session cost di chatroom.\n3. Agen mengirim pesan berbayar kedua (biaya Rp500).",
            "expected_desc": "1. Setelah pesan pertama terkirim, running cost langsung ter-update menjadi 'Rp500'.\n2. Setelah pesan kedua terkirim, running cost langsung ter-update menjadi 'Rp1.000' secara real-time tanpa refresh halaman.",
            "user_story": "User Story 4 - AC1"
        },
        {
            "title": "Owner/Supervisor mengatur Flat Non-Segmented Session Limit melalui quick-access di Chatroom",
            "decription": "Memastikan Owner/Supervisor dapat menyetel batas biaya flat per sesi percakapan secara langsung dari chatroom yang berlaku umum untuk semua percakapan.",
            "precondition": "1. Login sebagai Owner / Supervisor.\n2. Membuka halaman Chatroom.",
            "priority": "P2 - High",
            "step_desc": "1. Klik ikon pengaturan quick-access 'Session Spend Limit' di chatroom.\n2. Masukkan nilai batas nominal flat (misal: Rp5.000).\n3. Simpan pengaturan.",
            "expected_desc": "1. Pengaturan flat session limit berhasil tersimpan.\n2. Nilai limit Rp5.000 aktif sebagai batas default untuk sesi-sesi percakapan baru.",
            "user_story": "User Story 4 - AC2"
        },
        {
            "title": "Verifikasi perubahan Flat Session Limit hanya berlaku pada sesi percakapan baru (Sesi aktif tidak berubah)",
            "decription": "Memastikan jika Owner mengubah batas limit sesi di tengah jalan, sesi chat yang sedang berjalan tetap menggunakan limit lamanya hingga sesi tersebut berakhir.",
            "precondition": "1. Sesi Chat A sedang berjalan dengan limit awal Rp5.000 (running cost saat ini Rp3.000).\n2. Owner mengubah Flat Session Limit global menjadi Rp2.000.",
            "priority": "P2 - High",
            "step_desc": "1. Owner memperbarui flat session limit menjadi Rp2.000.\n2. Buka Sesi Chat A yang sedang berjalan.\n3. Kirim pesan di Sesi Chat A hingga running cost menjadi Rp3.500.\n4. Buka Sesi Chat B yang baru dimulai setelah perubahan limit.",
            "expected_desc": "1. Sesi Chat A TIDAK terblokir pada Rp3.500 karena limit aktifnya tetap Rp5.000.\n2. Sesi Chat B yang baru mengadopsi limit baru Rp2.000 dan akan terblokir jika mencapai Rp2.000.",
            "user_story": "User Story 4 - AC2 (Decision Rule)"
        },

        # ==========================================
        # TS08: Limit-Reached Enforcement & Cumulative Agent Top-Up (User Story 4)
        # ==========================================
        {
            "title": "Verifikasi percakapan ditandai 'limit-reached' dan pemblokiran pesan berbayar saat session cost mencapai limit",
            "decription": "Memastikan ketika running cost sesi mencapai atau melebihi flat session limit, sistem menandai percakapan sebagai limit-reached dan menonaktifkan pengiriman pesan berbayar.",
            "precondition": "1. Flat session limit aktif = Rp5.000.\n2. Running cost sesi chatroom telah mencapai Rp5.000.",
            "priority": "P2 - High",
            "step_desc": "1. Amati status percakapan di chatroom.\n2. CS mencoba mengetik dan mengirim pesan berbayar baru.",
            "expected_desc": "1. Percakapan diberi penanda/flag 'limit-reached'.\n2. Tombol kirim pesan berbayar dinonaktifkan (disabled) / diblokir.\n3. Tampil opsi bagi agen untuk menambah limit sesi.",
            "user_story": "User Story 4 - AC3"
        },
        {
            "title": "Agent/CS melakukan manual top-up limit sesi via preset dropdown (Rp2.000 / Rp5.000 / Rp10.000)",
            "decription": "Memastikan agen (termasuk role CS) dapat membuka blokir percakapan limit-reached dengan memilih preset nominal top-up.",
            "precondition": "1. Login sebagai CS Operasional.\n2. Membuka percakapan yang berstatus 'limit-reached'.",
            "priority": "P2 - High",
            "step_desc": "1. Buka dropdown penambahan limit di percakapan tersebut.\n2. Pilih nominal preset: 'Rp5.000'.\n3. Klik tombol konfirmasi penambahan limit.",
            "expected_desc": "1. Limit sesi bertambah sebesar Rp5.000 (total limit menjadi Rp10.000).\n2. Flag 'limit-reached' hilang dan percakapan unblocked.\n3. CS dapat kembali mengirimkan pesan berbayar.",
            "user_story": "User Story 4 - AC4"
        },
        {
            "title": "Verifikasi sifat kumulatif pada multiple manual top-up oleh Agent/CS dalam satu sesi",
            "decription": "Memastikan jika agen melakukan top-up beberapa kali berturut-turut dalam satu sesi, nilai limit sesi terakumulasi secara kumulatif.",
            "precondition": "1. Sesi limit awal: Rp5.000 (status limit-reached).\n2. CS melakukan top-up pertama sebesar Rp2.000.\n3. CS kemudian melakukan top-up kedua sebesar Rp5.000.",
            "priority": "P2 - High",
            "step_desc": "1. Pilih preset Rp2.000 dan konfirmasi (Limit menjadi Rp7.000).\n2. Pilih kembali preset Rp5.000 dan konfirmasi.\n3. Amati batas maksimal limit sesi yang tercatat.",
            "expected_desc": "1. Limit sesi bertambah secara kumulatif menjadi Rp12.000 (Rp5.000 awal + Rp2.000 + Rp5.000).\n2. Pengiriman pesan berbayar diizinkan hingga akumulasi biaya mencapai Rp12.000.",
            "user_story": "User Story 4 - AC4 (Decision Rule)"
        },
        {
            "title": "Verifikasi pencatatan Audit Log Top-up Sesi (Agent Identity, Timestamp, Amount)",
            "decription": "Memastikan setiap aksi top-up limit oleh agen tercatat lengkap pada audit log sistem.",
            "precondition": "1. CS `CS_ANI` melakukan top-up limit Rp5.000 pada percakapan dengan Customer ID `CUST_77`.",
            "priority": "P3 - Medium",
            "step_desc": "1. Eksekusi konfirmasi top-up oleh CS `CS_ANI`.\n2. Periksa audit log aktivitas chatroom di sistem.",
            "expected_desc": "1. Log aktivitas merekam:\n   - Identitas Agen (`CS_ANI`)\n   - Percakapan/Session ID terkait\n   - Nominal yang ditambahkan (Rp5.000)\n   - Timestamp waktu konfirmasi top-up.",
            "user_story": "User Story 4 - AC4"
        },
        {
            "title": "Verifikasi reset limit ke nilai default Flat Session Limit saat sesi baru Meta dimulai",
            "decription": "Memastikan saat window 24 jam Meta berakhir dan customer memulai sesi baru, nilai limit sesi kembali ke default flat limit tanpa membawa sisa top-up sesi sebelumnya.",
            "precondition": "1. Sesi percakapan hari kemarin memiliki akumulasi limit Rp12.000 karena top-up.\n2. Window 24 jam Meta berakhir (sesi tertutup).\n3. Customer mengirimkan pesan baru (memulai sesi Meta baru).",
            "priority": "P2 - High",
            "step_desc": "1. Customer mengirimkan inbound pesan baru setelah window sesi lama expired.\n2. Periksa limit dan running cost sesi yang baru terbentuk.",
            "expected_desc": "1. Sesi baru terinisiasi dengan running cost Rp0.\n2. Limit sesi otomatis ter-reset ke nilai default Flat Session Limit global (misal: Rp5.000).\n3. Akumulasi top-up dari sesi kemarin tidak terbawa ke sesi baru.",
            "user_story": "User Story 4 - AC1, AC2 (Session Lifecycle)"
        }
    ]

    for row_idx, tc in enumerate(test_cases, start=2):
        row_data = [
            "",                                  # project_id
            "",                                  # suite_id
            "",                                  # unique_id
            tc["title"],                         # title
            tc["decription"],                    # decription
            tc["precondition"],                 # precondition
            tc["priority"],                      # priority
            "",                                  # type
            "",                                  # status
            "",                                  # tags
            "",                                  # steps
            tc["step_desc"],                     # step_desc
            tc["expected_desc"],                 # expected_desc
            "FALSE",                             # is_automated
            tc["user_story"]                     # user_story
        ]
        ws.append(row_data)

        # Styling data rows
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            # Center align priority & boolean
            if col_idx in [7, 14]:
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

    # Set column widths
    col_widths = {
        "A": 12,  # project_id
        "B": 12,  # suite_id
        "C": 12,  # unique_id
        "D": 35,  # title
        "E": 40,  # decription
        "F": 35,  # precondition
        "G": 16,  # priority
        "H": 10,  # type
        "I": 10,  # status
        "J": 10,  # tags
        "K": 10,  # steps
        "L": 45,  # step_desc
        "E": 40,  # decription
        "M": 45,  # expected_desc
        "N": 14,  # is_automated
        "O": 28   # user_story
    }

    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    output_path = "/Users/eldofadlyady/Documents/eldoTest/summary-temp/Test_Cases_Meta_New_Pricing_2026.xlsx"
    wb.save(output_path)
    print(f"File excel test case berhasil dibuat di: {output_path}")

if __name__ == "__main__":
    generate_test_cases()
