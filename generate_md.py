import openpyxl

wb = openpyxl.load_workbook('result-testcase/Test_Cases_Meta_New_Pricing_2026.xlsx')
ws = wb['Test Cases']

def get_complexity(uid, title, step_desc):
    if 'ROOM' in uid or 'CHT' in uid or 'Handover' in title or 'concurrency' in title.lower() or 'reset' in title.lower():
        if 'top-up' in title.lower() or 'dropdown' in step_desc.lower():
            return 'Medium-Complexity'
        if 'inbound' in title.lower() or 'closing' in step_desc.lower():
            return 'High-Complexity'
        return 'Medium-Complexity'
    if 'BRD' in uid or 'EXEC' in uid or 'FEP' in title or 'Rollback' in title:
        if 'Rollback' in title or 'CTWA' in title or 'WTWA' in title:
            return 'High-Complexity'
        return 'Medium-Complexity'
    if 'PRC' in uid or 'DB' in uid:
        return 'Low-Complexity'
    if 'CRD' in uid or 'Monthly' in title or 'Download' in title:
        if 'Download' in title or 'ekspor' in title.lower():
            return 'Medium-Complexity'
        return 'Low-Complexity'
    if 'PRE' in uid or 'AUTH' in uid:
        return 'Low-Complexity'
    return 'Low-Complexity'

with open('result-testcase/test-cases-meta-new-pricing-2026.md', 'w', encoding='utf-8') as f:
    f.write('# Automation Test Cases: Everpro Chat Meta New Pricing 2026

')
    f.write('> Dokumen ini dihasilkan otomatis sebagai turunan dari Master Test Case Excel () dan disesuaikan untuk instruksi test automation (Playwright / Cypress / Robot Framework).

')
    
    current_suite = ''
    for r in range(2, ws.max_row + 1):
        suite_id = ws.cell(r, 2).value or 'GENERAL'
        uid = ws.cell(r, 3).value or ''
        title = ws.cell(r, 4).value or ''
        desc = ws.cell(r, 5).value or ''
        precon = ws.cell(r, 6).value or ''
        prio = ws.cell(r, 7).value or 'Normal'
        tc_type = ws.cell(r, 8).value or 'Functional'
        step_desc = ws.cell(r, 12).value or ''
        exp_desc = ws.cell(r, 13).value or ''
        ustory = ws.cell(r, 15).value or ''

        if suite_id != current_suite:
            current_suite = suite_id
            f.write(f'## Suite: {current_suite}

')

        comp = get_complexity(uid, title, step_desc)
        tag_prio = '[Auto-' + prio + ']'
        tag_comp = '[' + comp + ']'
        f.write(f'### {uid}: {title}  
')
        f.write(f'- **User Story**: {ustory}
')
        f.write(f'- **Tipe**: {tc_type}
')
        f.write(f'- **Deskripsi**: {desc}
')
        f.write('- **Preconditions**:
')
        for line in precon.split('
'):
            if line.strip():
                f.write(f'  {line.strip()}
')
        
        f.write('- **Test Steps & Assertions**:
')
        steps_list = [s.strip() for s in step_desc.split('
') if s.strip()]
        exps_list = [e.strip() for e in exp_desc.split('
') if e.strip()]
        
        for s in steps_list:
            f.write(f'  {s}
')
        f.write('  **Expected Results**:
')
        for e in exps_list:
            f.write(f'  - {e}
')
        
        f.write('
---

')
