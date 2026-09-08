import os
import sys
import re
import argparse
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="6", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_row(row, bg_colors=None, is_header=False, bold=False, font_size=9.5):
    for i, cell in enumerate(row.cells):
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        if bg_colors:
            bg = bg_colors[i] if isinstance(bg_colors, list) else bg_colors
            if bg and bg != "auto":
                set_cell_background(cell, bg)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(font_size)
                if bold or is_header:
                    r.font.bold = True

def add_heading_styled(doc, text, level=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level==2 else 10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 3:
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(30, 30, 30)
    return p

def add_body_p(doc, text, bold_prefix=None, is_bullet=False, indent_level=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if is_bullet:
        p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
        run_bullet = p.add_run("•  ")
        run_bullet.font.name = 'Arial'
        run_bullet.font.size = Pt(9.5)
        run_bullet.font.bold = True
    elif indent_level > 0:
        p.paragraph_format.left_indent = Inches(0.25 * indent_level)
        
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = 'Arial'
        r_prefix.font.size = Pt(9.5)
        r_prefix.font.bold = True
        
    r_text = p.add_run(text)
    r_text.font.name = 'Arial'
    r_text.font.size = Pt(9.5)
    return p

def create_table_from_parsed(doc, headers, rows_data):
    if not headers and not rows_data:
        return
    all_rows = [headers] + rows_data if headers else rows_data
    num_cols = max(len(r) for r in all_rows)
    # pad rows if necessary
    for r in all_rows:
        while len(r) < num_cols:
            r.append("")

    table = doc.add_table(rows=len(all_rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="000000", sz="6")

    # Pick header color: Cyan if looks like test scenario/decision, else Gray
    header_str = " ".join(all_rows[0]).lower()
    is_cyan = any(w in header_str for w in ["scenario", "test case", "rule", "fep", "decision", "kondisi", "skenario"])
    header_bg = "00FFFF" if is_cyan else "CCCCCC"

    for r_idx, row_data in enumerate(all_rows):
        row = table.rows[r_idx]
        is_hdr = (r_idx == 0)
        bg = [header_bg] * num_cols if is_hdr else None
        
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(val).strip()
            # alignment
            txt = str(val).strip()
            if txt in ["✅", "❌", "Pending", "Success", "Delivered", "-", "TRUE", "FALSE"]:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        format_row(row, bg_colors=bg, is_header=is_hdr, font_size=9)

    doc.add_paragraph() # spacing

def convert_md_to_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    in_table = False
    table_headers = []
    table_rows = []

    # Title extraction
    title_text = ""
    for line in lines:
        if line.strip().startswith("# "):
            title_text = line.strip()[2:].strip()
            break
    if not title_text:
        title_text = os.path.basename(md_path).replace(".md", "")

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run(f"[TIW] {title_text}")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(14)
    r_title.font.bold = True

    # Header metadata box (Standard Everpro TIW)
    feature_name = title_text.replace("Analisis PRD & Perancangan Test Matrix:", "").strip()
    meta_data = [
        ["Product/Feature Name", feature_name],
        ["Supporting Docs", "PRD, Figma, RBAC SSOT, Adhoc Docs"],
        ["Contributor", "Eldo Fadlyady (QA Engineer)"],
        ["Approver", "Rizky Nuredja, Elsa Vinietta"],
        ["Informed", "Everpro Core Team, Product Management, Engineering"]
    ]
    meta_tbl = doc.add_table(rows=len(meta_data), cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_tbl, color="000000", sz="6")
    for idx, (k, v) in enumerate(meta_data):
        row = meta_tbl.rows[idx]
        row.cells[0].text = k
        row.cells[1].text = v
        set_cell_margins(row.cells[0], 100, 100, 150, 150)
        set_cell_margins(row.cells[1], 100, 100, 150, 150)
        set_cell_background(row.cells[0], "CCCCCC")
        format_row(row, bg_colors=["CCCCCC", "auto"], is_header=False, font_size=9.5)
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.6)
    doc.add_paragraph()

    # Changelog
    add_heading_styled(doc, "Changelog", level=2)
    changelog_data = [
        ["Version", "Date", "Description", "Update By", "Color Mark"],
        ["1.0.0", "May 18, 2026", "Initial Document Analysis & Test Matrix", "Eldo Fadlyady", ""]
    ]
    create_table_from_parsed(doc, changelog_data[0], changelog_data[1:])

    # Approval
    add_heading_styled(doc, "Approval", level=2)
    approval_data = [
        ["Name", "Status", "Date", "Notes"],
        ["Rizky Nuredja", "Pending", "", ""],
        ["Elsa Vinietta", "Pending", "", ""],
        ["QA Lead / PM", "Pending", "", ""],
        ["Engineering Lead", "Pending", "", ""]
    ]
    create_table_from_parsed(doc, approval_data[0], approval_data[1:])

    # Parse content
    for line in lines:
        raw = line.rstrip("\n\r")
        s = raw.strip()

        # Check table
        if s.startswith("|") and s.endswith("|"):
            cells = [c.strip() for c in s[1:-1].split("|")]
            # ignore divider row |:---|:---:|
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            if not in_table:
                in_table = True
                table_headers = cells
                table_rows = []
            else:
                table_rows.append(cells)
            continue
        else:
            if in_table:
                create_table_from_parsed(doc, table_headers, table_rows)
                in_table = False
                table_headers = []
                table_rows = []

        if not s:
            continue

        if s.startswith("# "):
            continue # Already handled as document title
        elif s.startswith("## "):
            add_heading_styled(doc, s[3:].strip(), level=2)
        elif s.startswith("### "):
            add_heading_styled(doc, s[4:].strip(), level=3)
        elif s.startswith("---"):
            continue
        elif s.startswith("- ") or s.startswith("* "):
            content = s[2:].strip()
            # check bold prefix like **Tujuan Fitur**: ...
            m = re.match(r'^\*\*(.*?)\*\*:(.*)$', content)
            if m:
                add_body_p(doc, m.group(2), bold_prefix=m.group(1) + ":", is_bullet=True)
            else:
                add_body_p(doc, content, is_bullet=True)
        elif re.match(r'^\d+\.\s+', s):
            content = re.sub(r'^\d+\.\s+', '', s).strip()
            num = re.match(r'^\d+\.', s).group(0)
            m = re.match(r'^\*\*(.*?)\*\*:(.*)$', content)
            if m:
                add_body_p(doc, m.group(2), bold_prefix=f"{num} {m.group(1)}:", is_bullet=False)
            else:
                add_body_p(doc, content, bold_prefix=f"{num} ", is_bullet=False)
        else:
            add_body_p(doc, s)

    if in_table:
        create_table_from_parsed(doc, table_headers, table_rows)

    # Add Feedbacks & Review Log section at the end
    add_heading_styled(doc, "Feedbacks & Review Log", level=2)
    feedback_headers = ["Reviewer and Description", "Type", "Action"]
    feedback_rows = [
        ["", "", ""],
        ["", "", ""],
        ["", "", ""]
    ]
    create_table_from_parsed(doc, feedback_headers, feedback_rows)

    os.makedirs(os.path.dirname(os.path.abspath(docx_path)), exist_ok=True)
    doc.save(docx_path)
    print(f"Successfully converted MD to DOCX: {docx_path}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Convert Markdown QA Analysis to Standard DOCX")
    parser.add_argument("input_md", help="Path to input markdown file")
    parser.add_argument("output_docx", help="Path to output docx file")
    args = parser.parse_args()
    convert_md_to_docx(args.input_md, args.output_docx)
